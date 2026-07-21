# -*- coding: utf-8 -*-
"""
Genera el workbook 'Benchmark_Pacto_Primera_Infancia.xlsx' con las pestañas:
- Benchmark
- Backbone_Explicito_Extra
- Conceptos
- Dimensiones
- Fuentes
Todas las referencias son verificables (URLs públicas).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------- estilos ----------
H = Font(bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor="1F4E5F")
SUB_FILL = PatternFill("solid", fgColor="DCE6EA")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CAT_FILL = {
    "PRIORITARIO": PatternFill("solid", fgColor="E8F1D4"),
    "FUNCIONAL":   PatternFill("solid", fgColor="FCEFD4"),
    "ADYACENTE":   PatternFill("solid", fgColor="E3E0EF"),
}

def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = H
        cell.fill = HEAD_FILL
        cell.alignment = WRAP
        cell.border = BORDER
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def finish(ws, widths, header_row=1):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER

# =====================================================================
# 1) BENCHMARK
# =====================================================================
ws = wb.active
ws.title = "Benchmark"

headers = [
    "Categoría", "Caso", "País / Región", "Foco (edad / tema)",
    "¿Backbone explícito?", "Inicio", "Organización columna vertebral",
    "Escuela de Collective Impact",
    "D1 · Gobernanza distribuida", "D2 · Acción colectiva tangible",
    "D3 · Activación de miembros", "D4 · Conexión de la red",
    "D5 · Medición compartida", "D6 · Aprendizaje y adaptación colectiva",
    "Aprendizaje clave para el Pacto", "Fuente verificable (URL)",
]
ws.append(headers)

rows = [
    # ---- PRIORITARIOS ----
    ["PRIORITARIO", "Logan Together", "Australia (Logan, QLD)",
     "Bienestar infantil 0–8 años", "Sí — 'Backbone team' explícito", "2015",
     "Logan Together Campaign/Backbone Team (aloja Griffith University)",
     "Collective Impact 3.0 (versión 'evolucionada': incidencia en políticas + participación comunitaria + movimiento)",
     "El backbone comparte gobernanza con 'Focus Communities' locales; combina liderazgo comunitario y co-diseño ('community-led, place-based').",
     "Metas concretas 0–8: reducir la vulnerabilidad al desarrollo al promedio estatal; suburbios muestran mejoras en los 5 dominios del AEDC.",
     "Moviliza padres, servicios, gobierno y voluntariado; enfoque de 'citizen engagement' y co-diseño.",
     "Teje 'Focus Communities' + socios sectoriales bajo una hoja de ruta común (Roadmap).",
     "Usa el Australian Early Development Census (AEDC) como medición poblacional compartida; equipo de datos y evaluación.",
     "Marco de CI 'evolucionado' complementado con curso de vida, prevención, factores de riesgo/protección; aprendizaje con 'shared data + local wisdom'.",
     "Modelo place-based con backbone liviano + medición poblacional oficial (AEDC) como 'norte' compartido; muestra cómo pasar de datos a acción local.",
     "https://www.logantogether.org.au/  |  https://aifs.gov.au/resources/practice-guides/collective-impact-evidence-and-implications-practice"],

    ["PRIORITARIO", "Primero lo Primero", "Colombia",
     "Primera infancia (<5 años)", "Sí — se define como vehículo de 'Impacto Colectivo'", "2016 aprox.",
     "United Way Colombia (convocante/columna vertebral) con fundaciones aliadas",
     "Collective Impact 2.0 con rasgos 3.0 (mecanismo clásico de IC + enfoque territorial/descentralizado)",
     "Alianza de decisores públicos y privados; trabaja 'desde cada territorio' con perspectiva descentralizada.",
     "Compromiso concreto: construcción/operación de 13 centros aeioTU en 5 departamentos; ~5,929 niñas y niños con atención integral.",
     "Reúne 11 organizaciones públicas, 5 empresas privadas y 14 organizaciones sociales/multilaterales.",
     "Vincula sector público, privado y multilateral en un mismo mecanismo nacional.",
     "Mecanismo de impacto colectivo orientado a mejorar condiciones de calidad y medir avance del sistema de primera infancia.",
     "Cambio de sistema para mejorar la vida en la primera infancia y transformar índices de pobreza/desigualdad; aprendizaje territorial.",
     "Referente latinoamericano de IC en primera infancia con backbone tipo United Way; útil por contexto comparable a México (público-privado + territorial).",
     "https://primeroloprimero.co/es/impacto-colectivo/  |  https://fundacioncompartir.org/noticias/alianza-primero-lo-primero-juntos-por-primera-infancia"],

    # ---- FUNCIONALES ----
    ["FUNCIONAL", "ECDAN (Early Childhood Development Action Network)", "Global",
     "Desarrollo de la primera infancia y cuidadores", "No se autodenomina backbone; cumple funciones de columna vertebral global", "2016",
     "Red global creada por el Banco Mundial (con UNICEF); presidida por UNESCO desde 2022",
     "Cercano a CI 3.0 (paradigma de movimiento / red que alinea actores)",
     "Consejo/liderazgo ejecutivo multi-actor; conecta y alinea socios sin ser una sola organización rectora.",
     "Cataliza acción colectiva en metas comunes (p. ej. apoyo a la crianza, medición del desarrollo infantil).",
     "Conecta y activa +86 organizaciones fundadoras y socios regionales/globales.",
     "Conecta movimientos y actores intersectoriales alrededor de metas comunes.",
     "Impulsa medición de resultados del desarrollo infantil y rendición de cuentas; comparte herramientas y guías.",
     "Comparte investigación, reportes de aprendizaje y buenas prácticas mediante eventos y recursos descargables.",
     "Muestra cómo una red global ejerce funciones de backbone (alinear, gestionar conocimiento, incidir, medir) sin estructura jerárquica.",
     "https://ecdan.org/"],

    ["FUNCIONAL", "First Things First (FTF)", "EE. UU. (Arizona)",
     "Primera infancia, nacimiento–5 años", "No se llama 'backbone'; opera como columna vertebral estatal (agencia pública financiadora + coordinadora)", "2006",
     "First Things First (agencia estatal de primera infancia)",
     "Old-school / CI 2.0 (visión compartida + medición por indicadores + estructura formal; anterior al término 'Collective Impact')",
     "28 consejos regionales de voluntarios locales deciden el uso de fondos según necesidades de su comunidad.",
     "Invierte en programas probados: calidad de entornos de aprendizaje temprano, apoyo a familias, salud preventiva.",
     "Consejos integran padres, educadores, líderes empresariales, representantes tribales, profesionales de salud y filantropía.",
     "Articula el sistema estatal de primera infancia como 'socio clave' que coordina actores.",
     "Visión compartida de 'preparación para kínder' con indicadores comunes a nivel estatal/regional.",
     "Ajusta estrategias regionales según estudio de necesidades locales; decisiones basadas en datos.",
     "Modelo de backbone financiado con política pública estable y estructura regional; útil para pensar la relación Pacto–gobierno/estados.",
     "https://www.firstthingsfirst.org/what-we-do/how-we-work/"],

    # ---- ADYACENTES ----
    ["ADYACENTE", "AfECN (Africa Early Childhood Network)", "África (subsahariana)",
     "Desarrollo de la primera infancia (ECD)", "No — red/plataforma continental", "2015",
     "AfECN (organización sin fines de lucro registrada; secretaría en Nairobi)",
     "Movimiento/red con rasgos de CI 3.0 (construcción de campo y de comunidad)",
     "Plataforma que agrupa sociedad civil, academia, gobierno y sector privado a nivel nacional y regional.",
     "Cinco programas centrales orientados a mejores resultados para la niñez; incidencia, investigación, conocimiento.",
     "Activa 'comunidades' de miembros y redes nacionales en el continente.",
     "Fortalece alianzas y comparte experiencias/conocimiento en ECD en África.",
     "Genera investigación y conocimiento como bien común del campo (no un sistema de indicadores único).",
     "Comparte experiencias y aprendizajes entre países para influir en política y práctica.",
     "Prácticas para 'tejer red' continental: gobernanza de membresía, incidencia regional y gestión de conocimiento distribuido.",
     "https://afecn.org/"],

    ["ADYACENTE", "ARNEC (Asia-Pacific Regional Network for Early Childhood)", "Asia-Pacífico (47 países)",
     "Desarrollo de la primera infancia (ECD)", "No — red/plataforma regional", "2008",
     "ARNEC (secretaría regional, con sede en Singapur)",
     "Movimiento/red con rasgos de CI 3.0 (activación y conexión de miembros)",
     "Alianza intersectorial entre disciplinas, organizaciones, agencias e instituciones para avanzar la agenda de ECD.",
     "Impulsa inversión y agenda de ECD holística e inclusiva; conferencias anuales para movilizar el campo.",
     "Diseña entornos donde los miembros se sienten motivados y conectados para participar y contribuir.",
     "Plataforma de interacción y colaboración entre profesionales de ECD de sectores diversos.",
     "Equipa a los miembros con conocimiento actualizado para ser mejores voceros (advocacy).",
     "Facilita intercambio de conocimiento y aprendizaje entre 47 países.",
     "Prácticas de 'activación de miembros' y 'conexión de red' a gran escala; modelo de membresía diversa y representativa.",
     "https://arnec.net/"],
]
for r in rows:
    ws.append(r)

style_header(ws)
# color por categoría
for i, r in enumerate(rows, start=2):
    ws.cell(row=i, column=1).fill = CAT_FILL[r[0]]
    ws.cell(row=i, column=1).font = Font(bold=True)
finish(ws, [13, 26, 18, 20, 22, 8, 26, 26, 34, 34, 30, 28, 30, 34, 36, 42])

# =====================================================================
# 2) BACKBONE EXPLÍCITO — EJEMPLOS EXTRA (otros temas de niñez / otros temas)
# =====================================================================
ws2 = wb.create_sheet("Backbone_Explicito_Extra")
h2 = ["Caso", "País / Región", "Tema", "¿Backbone explícito?",
      "Organización columna vertebral", "Escuela de Collective Impact",
      "Por qué es un backbone explícito de referencia", "Fuente verificable (URL)"]
ws2.append(h2)
extra = [
    ["StriveTogether — Cradle to Career Network", "EE. UU. (nacional)",
     "Éxito educativo de la niñez, 'de la cuna a la carrera' (7 hitos, incluye preparación para kínder)",
     "Sí — exige 'dedicated and sufficient backbone support'",
     "StriveTogether + backbones locales en ~100 comunidades",
     "CI 2.0 disciplinado (Theory of Action™) evolucionando hacia equidad",
     "Red nacional que institucionaliza el rol de backbone y la infraestructura cívica basada en datos; referencia canónica de backbone.",
     "https://www.strivetogether.org/what-we-do/collective-impact/"],
    ["Franklin County Communities That Care Coalition", "EE. UU. (Massachusetts)",
     "Prevención de consumo de sustancias en jóvenes / bienestar juvenil",
     "Sí — la Healthy Youth Partnership actúa como 'backbone organization'",
     "Healthy Youth Partnership (HYP)",
     "Old-school / CI 2.0 (perfilada por SSIR en 2012 como caso de Collective Impact)",
     "Caso clásico documentado por FSG y el Collective Impact Forum: backbone que coordina agencias y logra cambios a nivel poblacional.",
     "https://collectiveimpactforum.org/resource/case-study-communities-that-care-coalition/"],
    ["Vibrant Communities Canada — Cities Reducing Poverty", "Canadá (nacional, ~50 comunidades)",
     "Reducción de pobreza",
     "Sí — cada ciudad se estructura con backbone y medición compartida",
     "Tamarack Institute (convocante nacional) + backbones locales",
     "Cuna del pensamiento CI 3.0 (Cabaj & Weaver, Tamarack)",
     "Origen empírico de 'Collective Impact 3.0'; útil para ver gobernanza multinivel y medición compartida ('Game Changer').",
     "https://collectiveimpactforum.org/resource/case-study-vibrant-communities/"],
]
for r in extra:
    ws2.append(r)
style_header(ws2)
finish(ws2, [34, 24, 34, 30, 30, 30, 40, 44])

# =====================================================================
# 3) CONCEPTOS — definiciones de Collective Impact (la 'lente')
# =====================================================================
ws3 = wb.create_sheet("Conceptos")
h3 = ["#", "Definición / Marco de Collective Impact", "Autores (año)", "Fuente / publicación",
      "Descripción breve", "Atributos, elementos y criterios que la componen",
      "Lente que aporta al Benchmark", "URL verificable"]
ws3.append(h3)
concept = [
    [1, "Collective Impact (definición fundacional)", "Kania & Kramer (2011)",
     "Stanford Social Innovation Review (SSIR)",
     "'El compromiso de un grupo de actores importantes de distintos sectores con una agenda común para resolver un problema social específico', mediante colaboración estructurada.",
     "5 condiciones: (1) Agenda común, (2) Medición compartida, (3) Actividades que se refuerzan mutuamente, (4) Comunicación continua, (5) Organización de soporte (backbone).",
     "Lente clásica/estructurada, intersectorial y con backbone como eje.",
     "https://ssir.org/articles/entry/collective_impact"],
    [2, "Channeling Change: Making Collective Impact Work", "Hanleybrown, Kania & Kramer (2012)",
     "Stanford Social Innovation Review (SSIR)",
     "Operacionaliza cómo se pone en marcha el IC: precondiciones y fases de implementación.",
     "Precondiciones (campeón influyente, recursos financieros adecuados, urgencia del tema) + 3 fases (iniciar la acción, organizar para el impacto, sostener la acción y el impacto).",
     "Lente de arranque/implementación: qué debe existir antes y cómo madura.",
     "https://ssir.org/articles/entry/channeling_change_making_collective_impact_work"],
    [3, "Understanding the Value of Backbone Organizations", "Turner, Merchant, Kania & Martin (2012)",
     "SSIR / FSG (serie de blog)",
     "Define qué es y qué hace la organización columna vertebral (backbone).",
     "6 funciones del backbone: (1) guiar visión y estrategia, (2) apoyar actividades alineadas, (3) establecer prácticas de medición compartida, (4) construir voluntad pública, (5) impulsar políticas, (6) movilizar financiamiento.",
     "Lente para clasificar 'backbone explícito' vs. 'funcional' en el Benchmark.",
     "https://ssir.org/articles/entry/understanding_the_value_of_backbone_organizations_in_collective_impact_2"],
    [4, "Collective Impact 3.0 (marco evolucionado)", "Cabaj & Weaver (2016)",
     "Tamarack Institute",
     "Reinterpreta el IC desde un paradigma de construcción de movimiento, tras ~10 años de práctica en reducción de pobreza en Canadá. Distingue fases 1.0 / 2.0 / 3.0.",
     "5 condiciones evolucionadas: Aspiración comunitaria (vs agenda común), Aprendizaje estratégico (vs medición compartida), Actividades de alto apalancamiento (vs actividades que se refuerzan), Compromiso comunitario inclusivo (vs comunicación continua), Liderazgo/'containers for change' (vs backbone).",
     "Lente de 'nueva escuela': movimiento, participación comunitaria, adaptación.",
     "https://www.tamarackcommunity.ca/articles/collective-impact-3.0-an-evolving-framework-for-community-change"],
    [5, "Centering Equity in Collective Impact (redefinición)", "Kania, Williams, Schmitz, Brady, Kramer & Splansky Juster (2022)",
     "Stanford Social Innovation Review (SSIR)",
     "Redefine el IC poniendo la equidad como prerrequisito: 'una red de miembros de la comunidad, organizaciones e instituciones que promueven la equidad aprendiendo juntos, alineando e integrando sus acciones para lograr un cambio a nivel de población y de sistema'.",
     "5 estrategias de equidad: (1) fundamentar en datos y contexto, (2) enfocarse en cambio de sistemas, (3) redistribuir el poder, (4) escuchar y actuar con la comunidad, (5) construir liderazgo y rendición de cuentas en equidad.",
     "Lente de equidad y 'red' como unidad (no solo backbone).",
     "https://ssir.org/articles/entry/centering_equity_in_collective_impact"],
    [6, "Collective Impact: evidence and implications for practice", "Australian Institute of Family Studies — AIFS (guía de práctica)",
     "AIFS (Gobierno de Australia)",
     "Síntesis de evidencia sobre IC para practicantes; base usada por iniciativas como Logan Together.",
     "Revisa las 5 condiciones, la evidencia de efectividad, condiciones habilitantes y limitaciones del enfoque.",
     "Lente crítica/basada en evidencia para evaluar casos (útil para columna 'escuela de CI').",
     "https://aifs.gov.au/resources/practice-guides/collective-impact-evidence-and-implications-practice"],
    [7, "Collective Impact Principles of Practice", "Collective Impact Forum / FSG",
     "Collective Impact Forum",
     "Ocho principios de práctica que complementan las 5 condiciones con el 'cómo' cotidiano.",
     "8 principios (entre ellos: diseñar con la comunidad y no para ella; priorizar equidad; incluir a la comunidad; reclutar el liderazgo adecuado; usar datos para aprender, mejorar y rendir cuentas; cultivar liderazgo distribuido).",
     "Lente operativa/de principios para juzgar la calidad de la práctica.",
     "https://collectiveimpactforum.org/wp-content/uploads/2021/12/Collective-Impact-Principles-of-Practice.pdf"],
]
for r in concept:
    ws3.append(r)
style_header(ws3)
for i in range(2, 2 + len(concept)):
    ws3.cell(row=i, column=1).alignment = CENTER
finish(ws3, [4, 34, 26, 24, 46, 52, 34, 46])

# =====================================================================
# 4) DIMENSIONES — las 6 dimensiones y su trazabilidad a las definiciones de CI
# =====================================================================
ws4 = wb.create_sheet("Dimensiones")
h4 = ["#", "Dimensión", "Definición operativa (la 'lente' de esta columna)",
      "Condición(es) de CI de la que deriva", "Origen / fuente principal",
      "Qué observar en cada caso (criterios)"]
ws4.append(h4)
dims = [
    [1, "Gobernanza distribuida",
     "Liderazgo y toma de decisiones compartidos entre múltiples actores (no concentrados en una sola organización); estructuras adaptativas y liderazgo distribuido.",
     "Reformula 'Backbone support' (2011) hacia liderazgo adaptativo.",
     "Kania & Kramer (2011); Turner et al. (2012); 'containers for change' de CI 3.0 (Cabaj & Weaver 2016); 'shift power' (2022).",
     "¿Quién decide? ¿Hay backbone único o gobernanza multinivel/compartida? ¿Roles y estructuras claras?"],
    [2, "Acción colectiva tangible",
     "Actividades coordinadas y mutuamente reforzadas que producen cambios concretos y verificables (no solo coordinación o diálogo).",
     "'Mutually reinforcing activities' (2011) → 'high-leverage activities' (CI 3.0).",
     "Kania & Kramer (2011); Cabaj & Weaver (2016).",
     "¿Qué se hizo en la realidad? ¿Hay entregables/resultados medibles y no solo mesas de trabajo?"],
    [3, "Activación de miembros",
     "Involucramiento activo y sostenido de integrantes y comunidad; membresía que participa, aporta y se apropia.",
     "'Continuous communication' + 'inclusive community engagement' (CI 3.0) + 'listen to & act with community' (2022).",
     "Cabaj & Weaver (2016); Kania et al. (2022).",
     "¿Los miembros participan o solo figuran? ¿Hay mecanismos de activación y contribución?"],
    [4, "Conexión de la red",
     "Tejido de relaciones, flujo de información y alineación entre nodos/organizaciones; la red como unidad de análisis.",
     "'Continuous communication' (2011) + la 'red' como definición (2022) + network weaving.",
     "Kania & Kramer (2011); redefinición de equidad (2022).",
     "¿Cómo se conectan los nodos? ¿Hay tejido/plataforma o solo relación bilateral con el centro?"],
    [5, "Medición compartida",
     "Indicadores y sistemas comunes para medir y reportar progreso; datos usados para aprender, no solo para reportar.",
     "'Shared measurement' (2011) → 'strategic learning' (CI 3.0); 'ground in data' (2022).",
     "Kania & Kramer (2011); Cabaj & Weaver (2016); Kania et al. (2022).",
     "¿Hay indicadores comunes? ¿Fuente de datos poblacional? ¿Se usan para decidir?"],
    [6, "Aprendizaje y adaptación colectiva",
     "Reflexión conjunta y ajuste de estrategia basado en evidencia y contexto emergente (aprendizaje emergente/estratégico).",
     "'Strategic learning' (CI 3.0) + carácter emergente/adaptativo de la comunicación continua.",
     "Cabaj & Weaver (2016); Principios de Práctica (CI Forum).",
     "¿La iniciativa aprende y cambia de rumbo? ¿Hay ciclos de reflexión-ajuste documentados?"],
]
for r in dims:
    ws4.append(r)
style_header(ws4)
for i in range(2, 2 + len(dims)):
    ws4.cell(row=i, column=1).alignment = CENTER
    ws4.cell(row=i, column=2).font = Font(bold=True)
finish(ws4, [4, 26, 48, 40, 44, 44])

# Nota metodológica bajo la tabla
note_row = ws4.max_row + 2
ws4.cell(row=note_row, column=1, value="Nota metodológica:").font = Font(bold=True)
note = ("Sí: las 6 dimensiones funcionan como un 'greatest hits' que sintetiza y traduce las distintas "
        "definiciones de Collective Impact. La base son las 5 condiciones de Kania & Kramer (2011), releídas "
        "con la lente de CI 3.0 (Cabaj & Weaver, 2016) y la redefinición con equidad (Kania et al., 2022). "
        "Dos decisiones son deliberadas y conviene explicitarlas: (a) el 'Backbone' se reencuadra como "
        "'Gobernanza distribuida' (un movimiento propio de CI 3.0); y (b) la 'Agenda común' no aparece como "
        "dimensión propia —queda implícita en Gobernanza y en Acción colectiva—, por lo que vale la pena decidir "
        "si se hace explícita o se documenta por qué se integró.")
ws4.cell(row=note_row + 1, column=1, value=note).alignment = WRAP
ws4.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=6)
ws4.row_dimensions[note_row + 1].height = 92

# =====================================================================
# 5) FUENTES
# =====================================================================
ws5 = wb.create_sheet("Fuentes")
ws5.append(["Referencia", "URL"])
sources = [
    ["Kania & Kramer (2011), 'Collective Impact', SSIR", "https://ssir.org/articles/entry/collective_impact"],
    ["Hanleybrown, Kania & Kramer (2012), 'Channeling Change', SSIR", "https://ssir.org/articles/entry/channeling_change_making_collective_impact_work"],
    ["Turner, Merchant, Kania & Martin (2012), 'Understanding the Value of Backbone Organizations', SSIR/FSG", "https://ssir.org/articles/entry/understanding_the_value_of_backbone_organizations_in_collective_impact_2"],
    ["Cabaj & Weaver (2016), 'Collective Impact 3.0', Tamarack Institute", "https://www.tamarackcommunity.ca/articles/collective-impact-3.0-an-evolving-framework-for-community-change"],
    ["Kania, Williams, Schmitz, Brady, Kramer & Splansky Juster (2022), 'Centering Equity in Collective Impact', SSIR", "https://ssir.org/articles/entry/centering_equity_in_collective_impact"],
    ["AIFS, 'Collective Impact: evidence and implications for practice'", "https://aifs.gov.au/resources/practice-guides/collective-impact-evidence-and-implications-practice"],
    ["Collective Impact Forum, 'Principles of Practice'", "https://collectiveimpactforum.org/wp-content/uploads/2021/12/Collective-Impact-Principles-of-Practice.pdf"],
    ["Pacto por la Primera Infancia (México)", "https://www.pactoprimerainfancia.org.mx/"],
    ["Logan Together (Australia)", "https://www.logantogether.org.au/"],
    ["Primero lo Primero (Colombia) — Impacto Colectivo", "https://primeroloprimero.co/es/impacto-colectivo/"],
    ["ECDAN — Early Childhood Development Action Network", "https://ecdan.org/"],
    ["First Things First (Arizona) — How We Work", "https://www.firstthingsfirst.org/what-we-do/how-we-work/"],
    ["AfECN — Africa Early Childhood Network", "https://afecn.org/"],
    ["ARNEC — Asia-Pacific Regional Network for Early Childhood", "https://arnec.net/"],
    ["StriveTogether — Collective Impact", "https://www.strivetogether.org/what-we-do/collective-impact/"],
    ["Franklin County Communities That Care — caso (CI Forum)", "https://collectiveimpactforum.org/resource/case-study-communities-that-care-coalition/"],
    ["Vibrant Communities — caso (CI Forum)", "https://collectiveimpactforum.org/resource/case-study-vibrant-communities/"],
]
for r in sources:
    ws5.append(r)
style_header(ws5)
finish(ws5, [66, 90])

wb.save("/home/user/One-/Benchmark_Pacto_Primera_Infancia.xlsx")
print("OK -> Benchmark_Pacto_Primera_Infancia.xlsx")
print("Sheets:", wb.sheetnames)
